#imports
import os
import openpyxl

def getFileNames(dir):
        """
        Goes into the storage directory and returns a list of all file names.
        """
        if not os.path.exists(dir):
            raise FileNotFoundError(f"The directory {dir} does not exist.")
        
        # Returns only files (filtering out any subfolders if they exist)
        return [f for f in os.listdir(dir) if os.path.isfile(os.path.join(dir, f))]

class ExcelColumnManager:
    def __init__(self, filename, storage_dir):
        """
        Initializes the manager with a parametric file name and target directory.
        """
        self.filename = filename
        self.storage_dir = storage_dir
        self.file_path = os.path.join(self.storage_dir, self.filename)        

    def getList(self, column_letter, sheet_name=None):
        """
        Extracts all data from a specific column and returns it as a list.
        :param column_letter: The column letter (e.g., 'A' for column a).
        :param sheet_name: Optional sheet name. Defaults to the active sheet.
        """
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(f"The file {self.file_path} does not exist.")

        wb = openpyxl.load_workbook(self.file_path)
        sheet = wb[sheet_name] if sheet_name else wb.active

        column_data = []
        # Iterate through cells in the specified column
        for cell in sheet[column_letter]:
            column_data.append(cell.value)
            
        wb.close()
        return column_data

    def overight_list(self, column_letter, data_list, sheet_name=None):
        """
        Clears the specified column and writes new data vertically from a Python list.
        :param column_letter: The column letter (e.g., 'A' for column a).
        :param data_list: A list of values to write into the column.
        :param sheet_name: Optional sheet name. Defaults to the active sheet.
        """
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(f"The file {self.file_path} does not exist.")

        wb = openpyxl.load_workbook(self.file_path)
        sheet = wb[sheet_name] if sheet_name else wb.active

        # Clear the existing column by setting cell values to None up to the max row
        max_row = sheet.max_row
        for row in range(1, max_row + 1):
            sheet[f"{column_letter}{row}"] = None

        # Write the new data list vertically down the column starting from row 1
        for row_index, value in enumerate(data_list, start=1):
            sheet[f"{column_letter}{row_index}"] = value

        wb.save(self.file_path)
        wb.close()

    def add2List(self, name):
        newlist = self.getList("A")
        newlist.append(name)
        self.overight_list("A", newlist)

    def removeFromList(self, name):
        oldlist = self.getList("A")
        newlist = []
        for i in oldlist:
            if i != name:
                newlist.append(i)
        self.overight_list("A", newlist)

##interface---------------

